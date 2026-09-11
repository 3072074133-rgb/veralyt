from __future__ import annotations

import csv
import io
import logging
import re
import time
import uuid
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any, Iterable

import duckdb
import polars as pl
from charset_normalizer import from_bytes
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from .config import settings
from .models import DatasetColumn, DatasetInfo, DatasetRegion, UploadedFile
from .observability import bind_log_context, duration_ms, log_event
from .repository import repository
from .row_identity import INTERNAL_ROW_ID, initialize_row_sequence


ALLOWED_EXTENSIONS = {".xlsx", ".csv"}
logger = logging.getLogger(__name__)


class IngestionError(ValueError):
    def __init__(self, message: str, *, code: str = "ingestion_failed", http_status: int = 422) -> None:
        super().__init__(message)
        self.code = code
        self.http_status = http_status


@dataclass(frozen=True)
class ParsedRegion:
    sheet_name: str
    region_index: int
    display_name: str
    source_range: str
    header_start_row: int
    header_end_row: int
    data_start_row: int
    data_end_row: int
    frame: pl.DataFrame


@dataclass(frozen=True)
class WorkbookReadResult:
    regions: list[ParsedRegion]
    detected_sheet_count: int
    imported_sheet_count: int
    skipped_sheet_count: int

    @property
    def frames(self) -> list[tuple[str, pl.DataFrame]]:
        """Compatibility view used by existing callers and tests."""
        return [(region.display_name, region.frame) for region in self.regions]


@dataclass
class IngestionBudget:
    nonempty_cells: int = 0


def task_dir(task_id: str) -> Path:
    root = settings.data_dir.resolve()
    target = (root / "tasks" / task_id).resolve()
    if root not in target.parents:
        raise IngestionError("任务目录不合法")
    target.mkdir(parents=True, exist_ok=True)
    return target


def validate_file(path: Path, original_name: str, size: int) -> None:
    extension = Path(original_name).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise IngestionError("仅支持 .xlsx 和 .csv 文件", code="unsupported_type", http_status=415)
    if size <= 0:
        raise IngestionError("文件为空")
    if size > settings.max_file_size:
        limit_mb = settings.max_file_size // (1024 * 1024)
        raise IngestionError(f"单个文件不能超过 {limit_mb} MB", code="file_too_large", http_status=413)
    with path.open("rb") as stream:
        prefix = stream.read(8)
    if extension == ".xlsx":
        if not prefix.startswith(b"PK") or not zipfile.is_zipfile(path):
            raise IngestionError("文件内容不是有效的 XLSX 工作簿")
        with zipfile.ZipFile(path) as archive:
            names = set(archive.namelist())
            total_uncompressed = sum(item.file_size for item in archive.infolist())
            total_compressed = sum(max(item.compress_size, 1) for item in archive.infolist())
            if total_uncompressed > settings.max_xlsx_uncompressed_size:
                raise IngestionError(
                    "XLSX 解压后内容超过安全限制",
                    code="xlsx_uncompressed_too_large",
                    http_status=413,
                )
            if total_uncompressed / max(total_compressed, 1) > settings.max_xlsx_compression_ratio:
                raise IngestionError(
                    "XLSX 压缩比异常，已拒绝处理",
                    code="xlsx_compression_ratio_exceeded",
                    http_status=413,
                )
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            raise IngestionError("XLSX 文件缺少必要的工作簿结构")
    if extension == ".csv" and b"\x00" in prefix:
        raise IngestionError("CSV 文件包含无效的二进制内容")


def ingest_file(
    task_id: str,
    file_record: UploadedFile,
    stored_path: Path,
    *,
    publish_to_library: bool = False,
) -> list[DatasetInfo]:
    started_at = time.perf_counter()
    with bind_log_context(task_id=task_id, file_id=file_record.id):
        log_event(logger, "ingestion.started", file_type=stored_path.suffix.lower(), size=file_record.size)
        try:
            validate_file(stored_path, file_record.original_name, file_record.size)
            if stored_path.suffix.lower() == ".xlsx":
                workbook_result = _read_xlsx(stored_path)
                regions = workbook_result.regions
                file_record.detected_sheet_count = workbook_result.detected_sheet_count
                file_record.skipped_sheet_count = workbook_result.skipped_sheet_count
                imported_sheet_count = workbook_result.imported_sheet_count
            else:
                frame = _read_csv(stored_path)
                regions = [
                    ParsedRegion(
                        sheet_name=Path(file_record.original_name).stem,
                        region_index=1,
                        display_name=Path(file_record.original_name).stem,
                        source_range=f"A1:{get_column_letter(max(frame.width, 1))}{frame.height + 1}",
                        header_start_row=1,
                        header_end_row=1,
                        data_start_row=2,
                        data_end_row=frame.height + 1,
                        frame=frame,
                    )
                ]
                file_record.detected_sheet_count = 1
                file_record.skipped_sheet_count = 0
                imported_sheet_count = 1
        except Exception:
            logger.exception(
                "ingestion.failed",
                extra={"event_fields": {"event": "ingestion.failed", "duration_ms": duration_ms(started_at)}},
            )
            raise
    if not regions:
        raise IngestionError("文件中没有可分析的数据表")
    if any(region.frame.height > settings.max_rows_per_sheet for region in regions):
        raise IngestionError("数据区域行数超过安全限制", code="dataset_rows_exceeded", http_status=413)

    output_dir = task_dir(task_id) / "work"
    output_dir.mkdir(exist_ok=True)
    db_path = output_dir / "analysis.duckdb"
    datasets: list[DatasetInfo] = []
    parquet_paths: list[Path] = []
    created_tables: list[str] = []
    try:
        with duckdb.connect(str(db_path)) as connection:
            connection.execute("SET memory_limit='1GB'")
            connection.execute("SET threads=4")
            connection.execute("BEGIN TRANSACTION")
            for index, region in enumerate(regions):
                frame = _rename_reserved_row_id(region.frame)
                if frame.height == 0 or frame.width == 0:
                    continue
                dataset_id = str(uuid.uuid4())
                table_name = f"data_{file_record.id.replace('-', '')[:10]}_{index + 1}"
                parquet_path = output_dir / f"{dataset_id}.parquet"
                parquet_paths.append(parquet_path)
                created_tables.append(table_name)
                stored_frame = frame.with_row_index(INTERNAL_ROW_ID, offset=1).with_columns(
                    pl.col(INTERNAL_ROW_ID).cast(pl.Int64)
                )
                stored_frame.write_parquet(parquet_path)
                connection.execute(
                    f'CREATE OR REPLACE TABLE "{table_name}" AS SELECT * FROM read_parquet(?)',
                    [str(parquet_path)],
                )
                initialize_row_sequence(connection, dataset_id, table_name)
                columns = [
                    DatasetColumn(
                        name=name,
                        display_name=name,
                        data_type=str(frame.schema[name]),
                        null_count=frame[name].null_count(),
                        sample_values=(samples := [
                            _json_value(value) for value in frame[name].drop_nulls().head(3).to_list()
                        ]),
                        **_infer_column_semantics(name, str(frame.schema[name]), samples),
                    )
                    for name in frame.columns
                ]
                datasets.append(
                    DatasetInfo(
                        id=dataset_id,
                        file_id=file_record.id,
                        table_name=table_name,
                        display_name=region.display_name,
                        row_count=frame.height,
                        columns=columns,
                        source_region=DatasetRegion(
                            sheet_name=region.sheet_name,
                            region_index=region.region_index,
                            source_range=region.source_range,
                            header_start_row=region.header_start_row,
                            header_end_row=region.header_end_row,
                            data_start_row=region.data_start_row,
                            data_end_row=region.data_end_row,
                        ),
                    )
                )
            if not datasets:
                raise IngestionError("文件中的工作表均为空")
            connection.execute("COMMIT")
        file_record.status = "ready"
        file_record.sheet_count = imported_sheet_count
        file_record.row_count = sum(item.row_count for item in datasets)
        repository.publish_ingestion(
            task_id, file_record, datasets, publish_to_library=publish_to_library
        )
    except Exception:
        if db_path.exists() and created_tables:
            try:
                with duckdb.connect(str(db_path)) as cleanup:
                    for table_name in created_tables:
                        cleanup.execute(f'DROP TABLE IF EXISTS "{table_name}"')
            except Exception:
                pass
        for parquet_path in parquet_paths:
            parquet_path.unlink(missing_ok=True)
        raise
    with bind_log_context(task_id=task_id, file_id=file_record.id):
        log_event(
            logger,
            "ingestion.completed",
            datasets=len(datasets),
            rows=sum(item.row_count for item in datasets),
            duration_ms=duration_ms(started_at),
        )
    return datasets


def _rename_reserved_row_id(frame: pl.DataFrame) -> pl.DataFrame:
    if INTERNAL_ROW_ID not in frame.columns:
        return frame
    candidate = f"原始{INTERNAL_ROW_ID}"
    suffix = 2
    while candidate in frame.columns:
        candidate = f"原始{INTERNAL_ROW_ID}_{suffix}"
        suffix += 1
    return frame.rename({INTERNAL_ROW_ID: candidate})


def _read_xlsx(path: Path) -> WorkbookReadResult:
    table_ranges = _native_table_ranges(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    detected_sheet_count = len(workbook.sheetnames)
    if detected_sheet_count > settings.max_sheets:
        workbook.close()
        raise IngestionError(
            f"检测到 {detected_sheet_count} 个工作表，当前最多支持 {settings.max_sheets} 个"
        )
    regions: list[ParsedRegion] = []
    imported_sheet_count = 0
    skipped_sheet_count = 0
    budget = IngestionBudget()
    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                skipped_sheet_count += 1
                continue
            sheet_regions = _read_sheet_regions(worksheet, budget, table_ranges.get(worksheet.title, []))
            if not sheet_regions:
                skipped_sheet_count += 1
                continue
            imported_sheet_count += 1
            region_count = len(sheet_regions)
            for region in sheet_regions:
                display_name = (
                    worksheet.title
                    if region_count == 1
                    else f"{worksheet.title} · 区域 {region.region_index}"
                )
                regions.append(
                    ParsedRegion(
                        sheet_name=region.sheet_name,
                        region_index=region.region_index,
                        display_name=display_name,
                        source_range=region.source_range,
                        header_start_row=region.header_start_row,
                        header_end_row=region.header_end_row,
                        data_start_row=region.data_start_row,
                        data_end_row=region.data_end_row,
                        frame=_coerce_frame(region.frame),
                    )
                )
    finally:
        workbook.close()
    return WorkbookReadResult(
        regions=regions,
        detected_sheet_count=detected_sheet_count,
        imported_sheet_count=imported_sheet_count,
        skipped_sheet_count=skipped_sheet_count,
    )


SparseRow = tuple[int, dict[int, Any]]


def _native_table_ranges(path: Path) -> dict[str, list[str]]:
    """Read table metadata without loading a second, non-streaming workbook."""
    import posixpath
    ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    rid = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    with zipfile.ZipFile(path) as archive:
        def relations(part: str) -> dict[str, str]:
            directory, name = posixpath.split(part)
            rel = f'{directory}/_rels/{name}.rels'
            if rel not in archive.namelist():
                return {}
            return {r.attrib['Id']: posixpath.normpath(posixpath.join(directory, r.attrib['Target']))
                    if not r.attrib['Target'].startswith('/') else r.attrib['Target'].lstrip('/')
                    for r in ET.fromstring(archive.read(rel)) if r.attrib.get('TargetMode') != 'External'}
        book_rels = relations('xl/workbook.xml')
        result = {}
        for sheet in ET.fromstring(archive.read('xl/workbook.xml')).findall('s:sheets/s:sheet', ns):
            part = book_rels[sheet.attrib[rid]]
            sheet_rels = relations(part)
            refs = []
            for target in sheet_rels.values():
                if not target.startswith('xl/tables/'):
                    continue
                metadata = ET.fromstring(archive.read(target))
                if metadata.attrib.get('headerRowCount', '1') != '0':
                    refs.append(metadata.attrib['ref'])
            result[sheet.attrib['name']] = refs
        return result


def _read_sheet_regions(
    worksheet: Any, budget: IngestionBudget | None = None, table_ranges: list[str] | None = None
) -> list[ParsedRegion]:
    """Read one worksheet as sparse rows and split it into independent tables."""
    budget = budget or IngestionBudget()
    vertical_bands: list[list[SparseRow]] = []
    current_band: list[SparseRow] = []
    blank_run = 0

    for row_number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if row_number > settings.max_rows_per_sheet + 1:
            raise IngestionError(
                f"工作表“{worksheet.title}”行数超过 {settings.max_rows_per_sheet} 行限制",
                code="sheet_rows_exceeded",
                http_status=413,
            )
        sparse = {
            column_number: value
            for column_number, value in enumerate(values, start=1)
            if _is_present(value)
        }
        if sparse and max(sparse) > settings.max_columns_per_sheet:
            raise IngestionError(
                f"工作表“{worksheet.title}”列数超过 {settings.max_columns_per_sheet} 列限制",
                code="sheet_columns_exceeded",
                http_status=413,
            )
        budget.nonempty_cells += len(sparse)
        if budget.nonempty_cells > settings.max_nonempty_cells:
            raise IngestionError(
                f"工作簿非空单元格超过 {settings.max_nonempty_cells} 个限制",
                code="nonempty_cells_exceeded",
                http_status=413,
            )
        if sparse:
            blank_run = 0
            current_band.append((row_number, sparse))
            continue
        blank_run += 1
        if blank_run >= settings.region_blank_row_gap and current_band:
            vertical_bands.append(current_band)
            current_band = []

    if current_band:
        vertical_bands.append(current_band)

    # Financial statement sections share one header across blank separator rows.
    financial_names = {'利润表', '资产负债表', '现金流量表', '费用明细', '应收账款', '应付账款'}
    if worksheet.title in financial_names:
        rows = [row for band in vertical_bands for row in band]
        header = next((range_boundaries(ref)[1] for ref in table_ranges or []), None)
        header = header or next((r for r, v in rows if len(v) >= 3 and
                       any(x in v.values() for x in ('本月金额', '月末余额', '本月发生额')) and
                       any(x in v.values() for x in ('项目', '费用项目', '客户名称', '供应商名称'))), None)
        if header is None:
            raise IngestionError(f'工作表“{worksheet.title}”未识别到可靠财务表头，请检查项目及金额列',
                                 code='financial_header_invalid')
        rows = [(r, v) for r, v in rows if r >= header]
        width = max(rows[0][1])
        region = _build_region(sheet_name=worksheet.title, region_index=1, rows=rows,
                               start_column=1, end_column=width, forced_header_row=header)
        if region:
            source_rows = rows[1:]
            labels = [str(v.get(1, '')).strip() for _, v in source_rows]
            kinds = [('note' if len(v) == 1 or label.startswith(('编制', '口径', '数据说明', '测算'))
                      else 'check' if '核对' in label else 'total' if re.search(r'合计|总计', label)
                      else 'detail') for label, (_, v) in zip(labels, source_rows)]
            frame = region.frame.with_columns(pl.Series('来源行号', [r for r, _ in source_rows]),
                                               pl.Series('报表行类型', kinds))
            return [ParsedRegion(**{**region.__dict__, 'frame': frame})]

    result: list[ParsedRegion] = []
    for rows in vertical_bands:
        for start_column, end_column in _horizontal_regions(rows):
            region = _build_region(
                sheet_name=worksheet.title,
                region_index=len(result) + 1,
                rows=rows,
                start_column=start_column,
                end_column=end_column,
                forced_header_row=next((bounds[1] for ref in table_ranges or []
                                        if (bounds := range_boundaries(ref))[0] == start_column
                                        and bounds[2] == end_column
                                        and any(r == bounds[1] for r, _ in rows)), None),
            )
            if region is not None:
                result.append(region)
    return result


def _horizontal_regions(rows: list[SparseRow]) -> list[tuple[int, int]]:
    occupied_columns = sorted({column for _, values in rows for column in values})
    if not occupied_columns:
        return []
    result: list[tuple[int, int]] = []
    start = occupied_columns[0]
    previous = start
    for column in occupied_columns[1:]:
        blank_columns = column - previous - 1
        if blank_columns >= settings.region_blank_column_gap:
            result.append((start, previous))
            start = column
        previous = column
    result.append((start, previous))
    return result


def _build_region(
    *,
    sheet_name: str,
    region_index: int,
    rows: list[SparseRow],
    start_column: int,
    end_column: int,
    forced_header_row: int | None = None,
) -> ParsedRegion | None:
    region_rows = [
        (row_number, {column: value for column, value in values.items() if start_column <= column <= end_column})
        for row_number, values in rows
    ]
    region_rows = [(row_number, values) for row_number, values in region_rows if values]
    if len(region_rows) < 2:
        return None

    header_end_position = (next((i for i, (r, _) in enumerate(region_rows) if r == forced_header_row), None)
                           if forced_header_row is not None else
                           _find_header_position(region_rows, start_column, end_column))
    if header_end_position is None:
        return None
    header_start_position = header_end_position if forced_header_row is not None else _find_header_start(
        region_rows, header_end_position, start_column, end_column
    )
    headers = _expanded_headers(
        region_rows[header_start_position : header_end_position + 1],
        start_column,
        end_column,
    )
    data_rows = region_rows[header_end_position + 1 :]
    if not data_rows:
        return None

    dense_rows = [
        [values.get(column) for column in range(start_column, end_column + 1)]
        for _, values in data_rows
    ]
    frame = _frame_from_rows(dense_rows, headers)
    if frame.height == 0:
        return None

    first_source_row = region_rows[0][0]
    last_source_row = region_rows[-1][0]
    header_start_row = region_rows[header_start_position][0]
    header_end_row = region_rows[header_end_position][0]
    return ParsedRegion(
        sheet_name=sheet_name,
        region_index=region_index,
        display_name=sheet_name,
        source_range=(
            f"{get_column_letter(start_column)}{first_source_row}:"
            f"{get_column_letter(end_column)}{last_source_row}"
        ),
        header_start_row=header_start_row,
        header_end_row=header_end_row,
        data_start_row=data_rows[0][0],
        data_end_row=data_rows[-1][0],
        frame=frame,
    )


def _find_header_position(
    rows: list[SparseRow], start_column: int, end_column: int
) -> int | None:
    limit = min(len(rows), settings.header_search_rows)
    candidates = [
        (_header_score(rows, position, start_column, end_column), position)
        for position in range(limit)
        if position < len(rows) - 1
    ]
    if not candidates:
        return None
    score, position = max(candidates, key=lambda item: (item[0], -item[1]))
    if score == float("-inf"):
        return None
    while position + 2 < len(rows) and position < limit - 1:
        current_row_number = rows[position][0]
        next_row_number, next_values = rows[position + 1]
        next_profile = _row_profile(next_values, start_column, end_column)
        following_profile = _row_profile(rows[position + 2][1], start_column, end_column)
        is_leaf_header = (
            next_row_number == current_row_number + 1
            and len(rows[position][1]) < end_column - start_column + 1
            and next_profile["present"] >= 2
            # A sheet may place a second table beside a multi-level header.
            # Its numeric cells must not hide the textual leaf labels on the
            # left table (e.g. Inbound / Outbound / Extras).
            and next_profile["text"] >= 2
            and next_profile["text"] / next_profile["present"] >= 0.45
            and (next_profile["numeric"] == 0 or next_profile["numeric"] >= 2)
            and following_profile["numeric"] > 0
        )
        if not is_leaf_header:
            break
        position += 1
    return position


def _header_score(
    rows: list[SparseRow], position: int, start_column: int, end_column: int
) -> float:
    _, values = rows[position]
    profile = _row_profile(values, start_column, end_column)
    width = end_column - start_column + 1
    if profile["present"] == 0:
        return float("-inf")

    following = rows[position + 1 : position + 7]
    follower_profiles = [_row_profile(item, start_column, end_column) for _, item in following]
    numeric_followers = sum(item["numeric"] > 0 and item["present"] >= 2 for item in follower_profiles)
    structured_followers = sum(
        item["present"] >= max(1, min(profile["present"], 2)) for item in follower_profiles
    )
    if profile["present"] == 1 and (position != 0 or numeric_followers < 2):
        return float("-inf")
    text_ratio = profile["text"] / profile["present"]
    numeric_ratio = profile["numeric"] / profile["present"]
    header_terms = sum(
        bool(re.search(r"日期|月份|年度|部门|项目|名称|编号|收入|成本|费用|金额|数量|预算|实际|合计|占比", _header_text(value)))
        for value in values.values()
    )
    coverage = profile["present"] / width
    first_row_bonus = 3.0 if position == 0 else max(0.0, 1.2 - position * 0.2)
    return (
        text_ratio * 4.0
        - numeric_ratio * 2.0
        + coverage * 1.5
        + min(numeric_followers, 4) * 0.8
        + min(structured_followers, 4) * 0.35
        + min(header_terms, 4) * 0.45
        + _numeric_header_sequence_score(values, start_column, end_column)
        + first_row_bonus
        - position * 0.035
    )


def _numeric_header_sequence_score(
    values: dict[int, Any], start_column: int, end_column: int
) -> float:
    numeric_values = [
        value
        for column, value in sorted(values.items())
        if start_column <= column <= end_column
        and isinstance(value, (int, float, Decimal))
        and not isinstance(value, bool)
    ]
    if len(numeric_values) < 3:
        return 0.0
    as_float = [float(value) for value in numeric_values]
    consecutive = all(right - left == 1 for left, right in zip(as_float, as_float[1:]))
    month_like = min(as_float) >= 1 and max(as_float) <= 12
    return 3.5 if consecutive and month_like else 0.0


def _find_header_start(
    rows: list[SparseRow], header_end_position: int, start_column: int, end_column: int
) -> int:
    start = header_end_position
    width = end_column - start_column + 1
    while start > 0 and header_end_position - start < 2:
        previous_row_number, previous_values = rows[start - 1]
        current_row_number = rows[start][0]
        profile = _row_profile(previous_values, start_column, end_column)
        if previous_row_number + 1 != current_row_number:
            break
        if profile["present"] < 2 or profile["numeric"] / profile["present"] > 0.2:
            break
        has_group_shape = profile["present"] < width or any(
            re.search(r"预算|实际|同比|环比|本期|上期|累计|年度|季度", _header_text(value))
            for value in previous_values.values()
        )
        if not has_group_shape:
            break
        start -= 1
    return start


def _expanded_headers(
    header_rows: list[SparseRow], start_column: int, end_column: int
) -> list[str]:
    if len(header_rows) == 1:
        values = header_rows[0][1]
        return _unique_headers(values.get(column) for column in range(start_column, end_column + 1))

    layers: list[list[str]] = []
    for layer_index, (_, values) in enumerate(header_rows):
        current = ""
        layer: list[str] = []
        for column in range(start_column, end_column + 1):
            label = _header_text(values.get(column))
            # Numeric cells in a mixed second header row are usually values
            # from an adjacent table (for example 2601 or 2252), not labels.
            # Keep true textual leaf labels while allowing numeric headers in
            # a single-row table, where this filter is not applied.
            if layer_index == len(header_rows) - 1 and _looks_like_data_value(label):
                label = ""
            if label:
                current = label
            elif layer_index == len(header_rows) - 1:
                current = ""
            layer.append(label or current)
        layers.append(layer)

    combined: list[str] = []
    for column_offset in range(end_column - start_column + 1):
        parts: list[str] = []
        for layer in layers:
            value = layer[column_offset]
            if value and (not parts or parts[-1] != value):
                parts.append(value)
        combined.append(" / ".join(parts))
    return _unique_headers(combined)


def _looks_like_data_value(value: str) -> bool:
    """Return whether a leaf-header candidate is an unlabelled number."""
    if not value or not re.fullmatch(r"[-+]?\d+(?:\.\d+)?", value):
        return False
    # Keep short ordinal/month labels (1..12) available for ordinary headers.
    try:
        number = float(value)
    except ValueError:
        return False
    return number > 12 or len(value.lstrip("+-").split(".", 1)[0]) >= 3


def _row_profile(values: dict[int, Any], start_column: int, end_column: int) -> dict[str, int]:
    present_values = [
        value
        for column, value in values.items()
        if start_column <= column <= end_column and _is_present(value)
    ]
    numeric = sum(
        isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
        for value in present_values
    )
    text_values = sum(
        isinstance(value, (str, date, datetime)) or isinstance(value, bool)
        for value in present_values
    )
    return {"present": len(present_values), "numeric": numeric, "text": text_values}


def _frame_from_rows(data: list[list[Any]], headers: list[str]) -> pl.DataFrame:
    try:
        return pl.DataFrame(data, schema=headers, orient="row", strict=False)
    except Exception:
        string_rows = [[None if value is None else str(value) for value in row] for row in data]
        return pl.DataFrame(string_rows, schema=headers, orient="row", strict=False)


def _header_text(value: Any) -> str:
    if not _is_present(value):
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value).strip())[:80]


def _is_present(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _read_csv(path: Path) -> pl.DataFrame:
    raw = path.read_bytes()
    match = from_bytes(raw).best()
    if match is None:
        raise IngestionError("无法识别 CSV 文件编码")
    text = str(match)
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    _validate_csv_shape(text, delimiter)
    try:
        frame = pl.read_csv(
            io.StringIO(text), separator=delimiter, infer_schema_length=1000,
            try_parse_dates=True, ignore_errors=True, truncate_ragged_lines=True,
        )
        return frame.rename(_normalize_headers(frame.columns))
    except Exception as exc:
        raise IngestionError(f"CSV 解析失败：{exc}") from exc


def _validate_csv_shape(text: str, delimiter: str) -> None:
    nonempty_cells = 0
    try:
        for row_number, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
            if row_number > settings.max_rows_per_sheet + 1:
                raise IngestionError(
                    f"CSV 行数超过 {settings.max_rows_per_sheet} 行限制",
                    code="csv_rows_exceeded",
                    http_status=413,
                )
            if len(row) > settings.max_columns_per_sheet:
                raise IngestionError(
                    f"CSV 列数超过 {settings.max_columns_per_sheet} 列限制",
                    code="csv_columns_exceeded",
                    http_status=413,
                )
            if any(len(value) > settings.max_csv_field_chars for value in row):
                raise IngestionError(
                    "CSV 单元格文本超过安全限制",
                    code="csv_field_too_large",
                    http_status=413,
                )
            nonempty_cells += sum(bool(value.strip()) for value in row)
            if nonempty_cells > settings.max_nonempty_cells:
                raise IngestionError(
                    f"CSV 非空单元格超过 {settings.max_nonempty_cells} 个限制",
                    code="nonempty_cells_exceeded",
                    http_status=413,
                )
    except csv.Error as exc:
        raise IngestionError(f"CSV 结构无效：{exc}") from exc


def _rows_to_frame(rows: list[tuple[Any, ...]]) -> pl.DataFrame:
    sparse_rows: list[SparseRow] = []
    for row_number, row in enumerate(rows, start=1):
        values = {
            column_number: value
            for column_number, value in enumerate(row, start=1)
            if _is_present(value)
        }
        if values:
            sparse_rows.append((row_number, values))
    if not sparse_rows:
        return pl.DataFrame()
    start_column = min(column for _, values in sparse_rows for column in values)
    end_column = max(column for _, values in sparse_rows for column in values)
    region = _build_region(
        sheet_name="工作表",
        region_index=1,
        rows=sparse_rows,
        start_column=start_column,
        end_column=end_column,
    )
    if region is None:
        return pl.DataFrame()
    return region.frame


def _coerce_frame(frame: pl.DataFrame) -> pl.DataFrame:
    """Recover dates and numbers commonly stored as text in finance workbooks."""
    expressions: list[pl.Expr] = []
    for name, data_type in frame.schema.items():
        column = pl.col(name)
        if data_type != pl.String:
            expressions.append(column)
            continue
        values = frame[name].drop_nulls()
        if values.len() == 0:
            expressions.append(column)
            continue
        iso_dates = values.str.contains(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}$").sum()
        if iso_dates / values.len() >= 0.9:
            normalized_dates = column.str.replace_all("/", "-")
            expressions.append(normalized_dates.str.to_date("%Y-%m-%d", strict=False).alias(name))
            continue
        normalized = values.str.replace_all(",", "")
        number_values = normalized.cast(pl.Float64, strict=False)
        if number_values.drop_nulls().len() / values.len() >= 0.9:
            converted = column.str.replace_all(",", "").cast(pl.Float64, strict=False)
            if normalized.str.contains(r"\.\d+$").any():
                expressions.append(converted.alias(name))
            else:
                expressions.append(converted.cast(pl.Int64, strict=False).alias(name))
            continue
        expressions.append(column)
    return frame.select(expressions)


def _infer_column_semantics(name: str, data_type: str, samples: list[Any]) -> dict[str, Any]:
    """Infer conservative finance semantics without sending raw values to an LLM."""
    lowered = name.casefold().replace(" ", "")
    is_numeric = any(token in data_type.casefold() for token in ("int", "float", "decimal"))
    is_temporal = any(token in data_type.casefold() for token in ("date", "time")) or any(
        token in lowered for token in ("日期", "时间", "期间", "月份", "年度", "year", "month", "date")
    )
    is_identifier = any(
        token in lowered for token in ("编号", "编码", "单号", "凭证号", "客户id", "供应商id", "code")
    ) or lowered in {"id", "来源行号", "行次"}
    is_percentage = any(
        token in lowered for token in ("占比", "百分比", "比例", "比率", "percent", "ratio", "rate", "%")
    ) or lowered.endswith("率")
    is_amount = any(
        token in lowered
        for token in (
            "金额", "收入", "营收", "销售额", "成本", "费用", "利润", "余额", "预算", "实际",
            "应收", "应付", "回款", "税额", "amount", "revenue", "sales", "profit", "cost", "expense",
        )
    )

    if is_identifier:
        semantic_type, role, aggregation, confidence = "id", "identifier", "none", 0.95
    elif is_temporal:
        semantic_type, role, aggregation, confidence = "date", "dimension", "none", 0.95
    elif is_percentage:
        semantic_type, role, aggregation, confidence = "percentage", "measure", "average", 0.9
    elif is_amount and is_numeric:
        semantic_type, role, aggregation, confidence = "amount", "measure", "sum", 0.95
    elif is_numeric:
        semantic_type, role, aggregation, confidence = "metric", "measure", "sum", 0.7
    elif any(token in lowered for token in ("名称", "姓名", "name")):
        semantic_type, role, aggregation, confidence = "name", "dimension", "none", 0.85
    else:
        semantic_type, role, aggregation, confidence = "category", "dimension", "none", 0.65

    joined_samples = " ".join(str(value) for value in samples)
    currency = None
    if any(token in f"{name} {joined_samples}" for token in ("人民币", "¥", "￥", "CNY", "RMB")):
        currency = "CNY"
    elif "$" in joined_samples or "USD" in name.upper():
        currency = "USD"

    unit = "%" if semantic_type == "percentage" else None
    for candidate in ("万元", "亿元", "千元", "元"):
        if candidate in name:
            unit = candidate
            break
    return {
        "semantic_type": semantic_type,
        "role": role,
        "unit": unit,
        "currency": currency,
        "default_aggregation": aggregation,
        "semantic_confidence": confidence,
    }


def _normalize_headers(columns: Iterable[str]) -> dict[str, str]:
    normalized = _unique_headers(columns)
    return dict(zip(columns, normalized, strict=False))


def _unique_headers(values: Iterable[Any]) -> list[str]:
    headers: list[str] = []
    used: dict[str, int] = {}
    for index, value in enumerate(values):
        base = re.sub(r"\s+", " ", str(value).strip()) if value not in (None, "") else f"未命名列{index + 1}"
        base = base[:80]
        count = used.get(base, 0) + 1
        used[base] = count
        headers.append(base if count == 1 else f"{base}_{count}")
    return headers


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value
