from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook

from app.config import settings
from app.ingestion import (
    IngestionError,
    _coerce_frame,
    _read_csv,
    _read_xlsx,
    _rows_to_frame,
    validate_file,
)


def test_rows_to_frame_deduplicates_headers() -> None:
    frame = _rows_to_frame([("月份", "收入", "收入"), ("1月", 100, 90), ("2月", 120, 110)])
    assert frame.columns == ["月份", "收入", "收入_2"]
    assert frame.height == 2


def test_read_chinese_csv(tmp_path: Path) -> None:
    path = tmp_path / "data.csv"
    path.write_bytes("月份,收入\n1月,100\n2月,120\n".encode("gb18030"))
    frame = _read_csv(path)
    assert frame.height == 2
    assert "收入" in frame.columns


def test_excel_text_values_are_coerced() -> None:
    frame = pl.DataFrame({"日期": ["2025-01-31", "2025-02-28"], "收入": ["850,000", "920000"]})
    converted = _coerce_frame(frame)
    assert converted.schema["日期"] == pl.Date
    assert converted.schema["收入"] == pl.Int64


def test_xlsx_imports_visible_data_sheets_and_skips_hidden_sheet(tmp_path: Path) -> None:
    path = tmp_path / "financial-report.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "汇总"
    first.append(["月份", "收入"])
    first.append(["1月", 100])
    for index in range(2, 19):
        sheet = workbook.create_sheet(f"部门{index}")
        sheet.append(["月份", "收入"])
        sheet.append(["1月", index * 100])
    hidden = workbook.create_sheet("辅助表")
    hidden.append(["键", "值"])
    hidden.append(["a", 1])
    hidden.sheet_state = "hidden"
    workbook.save(path)

    result = _read_xlsx(path)

    assert result.detected_sheet_count == 19
    assert len(result.frames) == 18
    assert result.skipped_sheet_count == 1
    assert "辅助表" not in {name for name, _ in result.frames}


def test_xlsx_skips_empty_visible_sheet(tmp_path: Path) -> None:
    path = tmp_path / "with-empty-sheet.xlsx"
    workbook = Workbook()
    workbook.active.append(["月份", "收入"])
    workbook.active.append(["1月", 100])
    workbook.create_sheet("空白表")
    workbook.save(path)

    result = _read_xlsx(path)

    assert result.detected_sheet_count == 2
    assert len(result.frames) == 1
    assert result.skipped_sheet_count == 1


def test_xlsx_rejects_more_than_configured_sheet_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "too-many-sheets.xlsx"
    workbook = Workbook()
    for index in range(3):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.append(["月份", "收入"])
        sheet.append(["1月", 100])
    workbook.save(path)
    monkeypatch.setattr(settings, "max_sheets", 2)

    with pytest.raises(IngestionError, match="检测到 3 个工作表"):
        _read_xlsx(path)


def test_xlsx_expands_merged_multi_level_headers(tmp_path: Path) -> None:
    path = tmp_path / "multi-level.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经营表"
    sheet.append(["部门", "收入", None, "成本", None])
    sheet.append([None, "预算", "实际", "预算", "实际"])
    sheet.append(["销售", 100, 120, 60, 70])
    sheet.append(["运营", 80, 75, 50, 48])
    sheet.merge_cells("B1:C1")
    sheet.merge_cells("D1:E1")
    workbook.save(path)

    result = _read_xlsx(path)

    assert len(result.regions) == 1
    region = result.regions[0]
    assert region.frame.columns == ["部门", "收入 / 预算", "收入 / 实际", "成本 / 预算", "成本 / 实际"]
    assert region.frame.height == 2
    assert region.header_start_row == 1
    assert region.header_end_row == 2
    assert region.source_range == "A1:E4"


def test_xlsx_keeps_leaf_headers_when_adjacent_table_has_numeric_cells(tmp_path: Path) -> None:
    """A side-by-side table must not turn the first table's fields unnamed."""
    path = tmp_path / "adjacent-tables.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Group", "Kimmy", None, None, None, "date", "Eva", "Lydia"])
    sheet.append(["Month", "Inbound", "Outbound", "Extras", "Sum K", None, 2252, 6712])
    sheet.append(["JAN", 10, 20, 3, 33, "2601", 2252, 6712])
    sheet.append(["FEB", 12, 18, 4, 34, "2602", 2828, 6444])
    workbook.save(path)

    result = _read_xlsx(path)

    region = result.regions[0]
    assert region.header_start_row == 1
    assert region.header_end_row == 2
    assert region.frame.columns[:5] == [
        "Group / Month",
        "Kimmy / Inbound",
        "Kimmy / Outbound",
        "Kimmy / Extras",
        "Kimmy / Sum K",
    ]
    assert not any(name.startswith("未命名列") for name in region.frame.columns)


def test_xlsx_splits_horizontal_and_vertical_data_regions(tmp_path: Path) -> None:
    path = tmp_path / "regions.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "混合页面"
    sheet.append(["月份", "收入", None, None, "部门", "成本"])
    sheet.append(["1月", 100, None, None, "销售", 30])
    sheet.append(["2月", 120, None, None, "运营", 25])
    sheet.append([])
    sheet.append([])
    sheet.append(["项目", "金额"])
    sheet.append(["差旅", 12])
    sheet.append(["办公", 8])
    workbook.save(path)

    result = _read_xlsx(path)

    assert result.imported_sheet_count == 1
    assert len(result.regions) == 3
    assert [region.source_range for region in result.regions] == ["A1:B3", "E1:F3", "A6:B8"]
    assert [region.display_name for region in result.regions] == [
        "混合页面 · 区域 1",
        "混合页面 · 区域 2",
        "混合页面 · 区域 3",
    ]


def test_xlsx_finds_header_after_intro_rows_within_search_window(tmp_path: Path) -> None:
    path = tmp_path / "intro.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row_number in range(1, 31):
        sheet.cell(row=row_number, column=1, value=f"说明 {row_number}")
    sheet.append(["日期", "部门", "收入"])
    sheet.append(["2026-01-01", "销售", 100])
    sheet.append(["2026-02-01", "运营", 120])
    workbook.save(path)

    result = _read_xlsx(path)

    region = result.regions[0]
    assert region.frame.columns == ["日期", "部门", "收入"]
    assert region.header_start_row == 31
    assert region.data_start_row == 32
    assert region.frame.height == 2


def test_xlsx_rejects_uncompressed_content_over_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "large-uncompressed.xlsx"
    workbook = Workbook()
    workbook.active.append(["字段", "值"])
    workbook.active.append(["A", "x" * 2000])
    workbook.save(path)
    monkeypatch.setattr(settings, "max_xlsx_uncompressed_size", 128)

    with pytest.raises(IngestionError, match="解压后内容超过") as caught:
        validate_file(path, path.name, path.stat().st_size)

    assert caught.value.code == "xlsx_uncompressed_too_large"


def test_xlsx_rejects_nonempty_cell_limit(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "too-many-cells.xlsx"
    workbook = Workbook()
    workbook.active.append(["月份", "收入"])
    workbook.active.append(["1月", 100])
    workbook.save(path)
    monkeypatch.setattr(settings, "max_nonempty_cells", 3)

    with pytest.raises(IngestionError, match="非空单元格超过") as caught:
        _read_xlsx(path)

    assert caught.value.code == "nonempty_cells_exceeded"


def test_csv_rejects_row_column_and_field_limits(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "limited.csv"
    path.write_text("a,b\n1,secret-value\n2,ok\n", encoding="utf-8")

    monkeypatch.setattr(settings, "max_rows_per_sheet", 1)
    with pytest.raises(IngestionError, match="CSV 行数超过"):
        _read_csv(path)

    monkeypatch.setattr(settings, "max_rows_per_sheet", 10)
    monkeypatch.setattr(settings, "max_columns_per_sheet", 1)
    with pytest.raises(IngestionError, match="CSV 列数超过"):
        _read_csv(path)

    monkeypatch.setattr(settings, "max_columns_per_sheet", 10)
    monkeypatch.setattr(settings, "max_csv_field_chars", 5)
    with pytest.raises(IngestionError, match="单元格文本超过"):
        _read_csv(path)
