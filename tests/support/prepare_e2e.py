from __future__ import annotations

import json

from openpyxl import Workbook

from app.analysis_tools import execute_sql
from app.config import settings
from app.ingestion import ingest_file, task_dir
from app.models import (
    AnalysisDraft,
    ChartSeries,
    ChartSpec,
    DatasetInfo,
    EvidencePointer,
    Finding,
    Metric,
    UploadedFile,
)
from app.repository import repository


def main() -> None:
    repository.initialize()
    task_id = repository.create_task()
    upload_dir = task_dir(task_id) / "uploads"
    upload_dir.mkdir(exist_ok=True)
    csv_path = upload_dir / "e2e-finance.csv"
    # Include a UTF-8 BOM so encoding detection is deterministic when this
    # helper is spawned by Node.js under different Windows code pages.
    csv_path.write_text("月份,收入\n2026-01,100\n2026-02,120\n2026-03,90\n", encoding="utf-8-sig")
    file_record = UploadedFile(
        id="e2e-finance-file",
        original_name="e2e-finance.csv",
        size=csv_path.stat().st_size,
        status="ready",
    )
    repository.add_file(task_id, file_record, csv_path.name)
    datasets = ingest_file(task_id, file_record, csv_path)
    repository.publish_data_revision(task_id)
    dataset: DatasetInfo = datasets[0]

    run_id = repository.queue_analysis(task_id, "按月份分析收入趋势")
    repository.mark_execution_running(run_id)
    result = execute_sql(
        task_id,
        datasets,
        f'SELECT "月份", SUM("收入") AS "收入" FROM "{dataset.table_name}" '
        'GROUP BY "月份" ORDER BY "月份" ASC',
        "月度收入",
        run_id,
    )
    evidence_id = result.evidence_ids[0]
    pointer = EvidencePointer(
        evidence_id=evidence_id,
        row_index=1,
        field="收入",
        raw_value="120",
    )
    draft = AnalysisDraft(
        title="月度收入分析",
        summary="2026年2月收入为120。",
        summary_evidence_refs=[evidence_id],
        summary_evidence_pointers=[pointer],
        metrics=[Metric(label="2月收入", value="120", evidence_refs=[evidence_id], evidence_pointers=[pointer])],
        findings=[Finding(title="2月收入", detail="2026年2月收入为120。", evidence_refs=[evidence_id], evidence_pointers=[pointer])],
        charts=[ChartSpec(
            id="e2e-chart",
            title="月度收入趋势",
            chart_type="line",
            dataset_ref=evidence_id,
            category_field="月份",
            series=[ChartSeries(name="收入", field="收入")],
        )],
    )
    repository.finish_execution(run_id, "completed", result=draft, activate=True)
    repository.add_artifact(
        task_id,
        run_id,
        "query",
        "月度收入汇总",
        {
            "tool_name": result.tool_name,
            "summary": result.summary,
            "row_count": len(result.rows),
            "arguments": result.arguments,
            "columns": list(result.rows[0]),
            "rows": result.rows,
            "warnings": [],
        },
        result.evidence_ids,
    )
    repository.add_artifact(
        task_id,
        run_id,
        "validation",
        "证据校验通过",
        {"passed": True, "issues": [], "checked_evidence_ids": result.evidence_ids},
        result.evidence_ids,
    )
    repository.add_artifact(
        task_id,
        run_id,
        "chart",
        "月度收入趋势",
        {"chart": draft.charts[0].model_dump(mode="json")},
        result.evidence_ids,
    )
    workbook_path = settings.data_dir / "e2e-multi-sheet.xlsx"
    workbook = Workbook()
    for index in range(18):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = f"部门{index + 1}"
        sheet.append(["月份", "收入"])
        sheet.append(["2026-01", (index + 1) * 100])
    hidden = workbook.create_sheet("辅助表")
    hidden.append(["键", "值"])
    hidden.append(["A", 1])
    hidden.sheet_state = "hidden"
    workbook.save(workbook_path)

    print(json.dumps({"task_id": task_id, "workbook_path": str(workbook_path)}))


if __name__ == "__main__":
    main()
